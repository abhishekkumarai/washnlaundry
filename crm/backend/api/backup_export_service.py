"""Comprehensive Export and Backup generation for all CRM sections.

Supports:
- Individual section export as CSV or JSON (Staff, Orders, Attendance, Payroll, Customers, Expenses, Credits, Services)
- Full Backup as multi-sheet Excel (.xlsx) workbook containing all sections
- Full Backup as unified structured JSON file
"""

import csv
import io
import json
from datetime import datetime, date, time
from django.utils import timezone

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    openpyxl = None

from .models import (
    Customer, GarmentCategory, GarmentItem, Order, OrderItem,
    Expense, Credit, Staff, Attendance, SalaryPayment, SalaryAdvance,
)
from .tenancy import get_current_tenant


def _tenant_filter(qs, request=None):
    """Filters queryset by current tenant if available."""
    shop = None
    if request and hasattr(request, 'shop') and request.shop:
        shop = request.shop
    else:
        shop = get_current_tenant()
    if shop and hasattr(qs.model, 'shop'):
        return qs.filter(shop=shop)
    return qs


def _format_dt(val):
    if val is None:
        return ''
    if isinstance(val, (datetime, date)):
        return val.strftime('%Y-%m-%d %H:%M:%S') if isinstance(val, datetime) else val.strftime('%Y-%m-%d')
    if isinstance(val, time):
        return val.strftime('%H:%M:%S')
    return str(val)


# ── SECTION DATA EXTRACTORS ───────────────────────────────────────────────────

def get_staff_data(request=None):
    headers = [
        'ID', 'Name', 'Role', 'Phone', 'Email',
        'Monthly Wage', 'Status', 'App Login', 'Start Date',
    ]
    qs = _tenant_filter(Staff.objects.all(), request).order_by('name')
    rows = []
    json_data = []
    for s in qs:
        row = [
            str(s.id),
            s.name,
            s.role,
            s.phone,
            s.email or '',
            float(s.monthly_wage),
            s.status,
            'Yes' if s.has_app_login else 'No',
            _format_dt(s.start_date),
        ]
        rows.append(row)
        json_data.append({
            'id': str(s.id),
            'name': s.name,
            'role': s.role,
            'phone': s.phone,
            'email': s.email or '',
            'monthly_wage': float(s.monthly_wage),
            'status': s.status,
            'has_app_login': s.has_app_login,
            'start_date': _format_dt(s.start_date),
        })
    return headers, rows, json_data


def get_orders_data(request=None):
    headers = [
        'Order Number', 'Customer Name', 'Customer Phone', 'Status',
        'Payment Status', 'Payment Method', 'Delivery Type', 'Source',
        'Subtotal', 'Delivery Charge', 'Discount', 'Total Amount',
        'Paid Amount', 'Due Amount', 'Created At', 'Scheduled Date',
        'Address', 'Notes',
    ]
    qs = _tenant_filter(Order.objects.all(), request).order_by('-created_at')
    rows = []
    json_data = []
    for o in qs:
        row = [
            o.order_number,
            o.customer_name,
            o.customer_phone,
            o.status,
            o.payment_status,
            o.payment_method or '',
            o.delivery_type,
            o.source,
            float(o.subtotal or 0.0),
            float(o.delivery_charge or 0.0),
            float(o.discount_amount or 0.0),
            float(o.total_amount or 0.0),
            float(o.paid_amount or 0.0),
            float(o.due_amount or 0.0),
            _format_dt(o.created_at),
            _format_dt(o.scheduled_date),
            o.address or '',
            o.notes or '',
        ]
        rows.append(row)
        json_data.append({
            'id': str(o.id),
            'order_number': o.order_number,
            'customer_name': o.customer_name,
            'customer_phone': o.customer_phone,
            'status': o.status,
            'payment_status': o.payment_status,
            'payment_method': o.payment_method,
            'delivery_type': o.delivery_type,
            'source': o.source,
            'subtotal': float(o.subtotal or 0.0),
            'delivery_charge': float(o.delivery_charge or 0.0),
            'discount_amount': float(o.discount_amount or 0.0),
            'total_amount': float(o.total_amount or 0.0),
            'paid_amount': float(o.paid_amount or 0.0),
            'due_amount': float(o.due_amount or 0.0),
            'created_at': _format_dt(o.created_at),
            'scheduled_date': _format_dt(o.scheduled_date),
            'address': o.address,
            'notes': o.notes,
        })
    return headers, rows, json_data


def get_attendance_data(request=None):
    headers = [
        'Staff Name', 'Date', 'Status', 'Check-in Time', 'Notes',
    ]
    qs = _tenant_filter(Attendance.objects.select_related('staff').all(), request).order_by('-date', 'staff__name')
    rows = []
    json_data = []
    for a in qs:
        staff_name = a.staff.name if a.staff else 'Unknown'
        row = [
            staff_name,
            _format_dt(a.date),
            a.status,
            _format_dt(a.check_in_time),
            a.notes or '',
        ]
        rows.append(row)
        json_data.append({
            'staff_id': str(a.staff_id) if a.staff_id else '',
            'staff_name': staff_name,
            'date': _format_dt(a.date),
            'status': a.status,
            'check_in_time': _format_dt(a.check_in_time),
            'notes': a.notes or '',
        })
    return headers, rows, json_data


def get_payroll_data(request=None):
    headers = [
        'Staff Name', 'Month', 'Amount', 'Paid On',
        'Payment Method', 'Type', 'Notes',
    ]
    rows = []
    json_data = []

    payments = _tenant_filter(SalaryPayment.objects.select_related('staff').all(), request).order_by('-month', '-paid_on')
    for p in payments:
        staff_name = p.staff.name if p.staff else 'Unknown'
        row = [
            staff_name,
            _format_dt(p.month),
            float(p.amount),
            _format_dt(p.paid_on),
            p.method,
            'Salary Payout',
            p.note or '',
        ]
        rows.append(row)
        json_data.append({
            'staff_id': str(p.staff_id) if p.staff_id else '',
            'staff_name': staff_name,
            'month': _format_dt(p.month),
            'amount': float(p.amount),
            'paid_on': _format_dt(p.paid_on),
            'method': p.method,
            'type': 'Salary Payout',
            'note': p.note or '',
        })

    advances = _tenant_filter(SalaryAdvance.objects.select_related('staff').all(), request).order_by('-month', '-paid_on')
    for adv in advances:
        staff_name = adv.staff.name if adv.staff else 'Unknown'
        row = [
            staff_name,
            _format_dt(adv.month),
            float(adv.amount),
            _format_dt(adv.paid_on),
            adv.method,
            'Salary Advance',
            adv.note or '',
        ]
        rows.append(row)
        json_data.append({
            'staff_id': str(adv.staff_id) if adv.staff_id else '',
            'staff_name': staff_name,
            'month': _format_dt(adv.month),
            'amount': float(adv.amount),
            'paid_on': _format_dt(adv.paid_on),
            'method': adv.method,
            'type': 'Salary Advance',
            'note': adv.note or '',
        })

    return headers, rows, json_data


def get_customers_data(request=None):
    headers = [
        'Name', 'Phone', 'Email', 'Address', 'Area',
        'Total Orders', 'Total Spent', 'Due Amount', 'Notes', 'Created At',
    ]
    qs = _tenant_filter(Customer.objects.all(), request).order_by('name')
    rows = []
    json_data = []
    for c in qs:
        row = [
            c.name,
            c.phone,
            c.email or '',
            c.address or '',
            c.area or '',
            c.total_orders,
            float(c.total_spent or 0.0),
            float(c.due_amount or 0.0),
            c.notes or '',
            _format_dt(c.created_at),
        ]
        rows.append(row)
        json_data.append({
            'id': str(c.id),
            'name': c.name,
            'phone': c.phone,
            'email': c.email or '',
            'address': c.address or '',
            'area': c.area or '',
            'total_orders': c.total_orders,
            'total_spent': float(c.total_spent or 0.0),
            'due_amount': float(c.due_amount or 0.0),
            'notes': c.notes or '',
            'created_at': _format_dt(c.created_at),
        })
    return headers, rows, json_data


def get_expenses_data(request=None):
    headers = [
        'Title', 'Category', 'Amount', 'Payment Method', 'Date', 'Notes',
    ]
    qs = _tenant_filter(Expense.objects.all(), request).order_by('-date')
    rows = []
    json_data = []
    for e in qs:
        row = [
            e.title,
            e.category,
            float(e.amount),
            e.payment_method,
            _format_dt(e.date),
            e.notes or '',
        ]
        rows.append(row)
        json_data.append({
            'id': str(e.id),
            'title': e.title,
            'category': e.category,
            'amount': float(e.amount),
            'payment_method': e.payment_method,
            'date': _format_dt(e.date),
            'notes': e.notes or '',
        })
    return headers, rows, json_data


def get_credits_data(request=None):
    headers = [
        'Title', 'Category', 'Amount', 'Payment Method', 'Date', 'Notes',
    ]
    qs = _tenant_filter(Credit.objects.select_related('category').all(), request).order_by('-date')
    rows = []
    json_data = []
    for cr in qs:
        cat_name = cr.category.name if cr.category else 'Other'
        row = [
            cr.title,
            cat_name,
            float(cr.amount),
            cr.payment_method,
            _format_dt(cr.date),
            cr.notes or '',
        ]
        rows.append(row)
        json_data.append({
            'id': str(cr.id),
            'title': cr.title,
            'category': cat_name,
            'amount': float(cr.amount),
            'payment_method': cr.payment_method,
            'date': _format_dt(cr.date),
            'notes': cr.notes or '',
        })
    return headers, rows, json_data


def get_services_data(request=None):
    headers = [
        'Category', 'Item Name', 'Price', 'Unit', 'Is Active', 'Display Order',
    ]
    qs = _tenant_filter(GarmentItem.objects.select_related('category').all(), request).order_by(
        'category__display_order', 'display_order', 'id'
    )
    rows = []
    json_data = []
    for item in qs:
        cat_name = item.category.name if item.category else 'General'
        row = [
            cat_name,
            item.name,
            float(item.price),
            item.get_unit_display(),
            'Yes' if item.is_active else 'No',
            item.display_order,
        ]
        rows.append(row)
        json_data.append({
            'id': item.id,
            'category': cat_name,
            'name': item.name,
            'price': float(item.price),
            'unit': item.get_unit_display(),
            'is_active': item.is_active,
            'display_order': item.display_order,
        })
    return headers, rows, json_data


SECTIONS = {
    'staff': ('Staff Roster', get_staff_data),
    'orders': ('Orders', get_orders_data),
    'attendance': ('Attendance', get_attendance_data),
    'payroll': ('Payroll & Advances', get_payroll_data),
    'customers': ('Customers', get_customers_data),
    'expenses': ('Expenses', get_expenses_data),
    'credits': ('Credits', get_credits_data),
    'services': ('Services & Catalog', get_services_data),
}


# ── CSV & JSON EXPORTERS ──────────────────────────────────────────────────────

def export_section_csv(section_name, request=None):
    """Exports a single section as a CSV string."""
    section_info = SECTIONS.get(section_name)
    if not section_info:
        raise ValueError(f"Unknown section: {section_name}")
    _, extractor = section_info
    headers, rows, _ = extractor(request)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for r in rows:
        writer.writerow(r)
    return output.getvalue()


def export_section_json(section_name, request=None):
    """Exports a single section as structured JSON serializable list."""
    section_info = SECTIONS.get(section_name)
    if not section_info:
        raise ValueError(f"Unknown section: {section_name}")
    _, extractor = section_info
    _, _, json_data = extractor(request)
    return json_data


def export_all_json(request=None):
    """Exports a full backup of all sections as a comprehensive JSON dict."""
    timestamp = timezone.now().strftime('%Y-%m-%d %H:%M:%S')
    backup = {
        'version': '1.0',
        'exported_at': timestamp,
        'sections': {},
    }
    for key, (label, extractor) in SECTIONS.items():
        _, _, json_data = extractor(request)
        backup['sections'][key] = {
            'label': label,
            'count': len(json_data),
            'records': json_data,
        }
    return backup


# ── FULL EXCEL WORKBOOK BACKUP ────────────────────────────────────────────────

def export_full_backup_xlsx(request=None):
    """Generates a complete multi-sheet Excel (.xlsx) workbook containing all sections."""
    if openpyxl is None:
        raise RuntimeError("openpyxl is not installed on the server.")

    wb = openpyxl.Workbook()
    # Remove default sheet later or reuse it
    default_sheet = wb.active

    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='182C4F', end_color='182C4F', fill_type='solid')
    header_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

    cell_font = Font(name='Segoe UI', size=10)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )

    # 1. Summary sheet first
    ws_summary = wb.create_sheet(title='Backup Summary', index=0)
    ws_summary.row_dimensions[1].height = 28
    ws_summary.append(['CRM Section', 'Records Count', 'Status'])
    for col_idx in range(1, 4):
        c = ws_summary.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = thin_border

    summary_row = 2
    for key, (label, extractor) in SECTIONS.items():
        headers, rows, _ = extractor(request)
        # Create worksheet for section
        sheet_title = label[:31]  # Excel max 31 chars
        ws = wb.create_sheet(title=sheet_title)

        # Header row
        ws.row_dimensions[1].height = 26
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # Data rows
        for row_idx, r in enumerate(rows, 2):
            ws.row_dimensions[row_idx].height = 20
            for col_idx, val in enumerate(r, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.font = cell_font
                c.border = thin_border
                if isinstance(val, (int, float)):
                    c.alignment = Alignment(horizontal='right')
                elif isinstance(val, str) and (val in ('Yes', 'No', 'ACTIVE', 'INACTIVE', 'PRESENT', 'ABSENT', 'HALF_DAY', 'LEAVE')):
                    c.alignment = Alignment(horizontal='center')

        # Auto column width
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max((len(str(cell.value or '')) for cell in col), default=10)
            ws.column_dimensions[col_letter].width = max(min(max_len + 4, 40), 12)

        # Add to summary row
        ws_summary.row_dimensions[summary_row].height = 20
        ws_summary.append([label, len(rows), 'Backed Up'])
        for col_idx in range(1, 4):
            c = ws_summary.cell(row=summary_row, column=col_idx)
            c.font = cell_font
            c.border = thin_border
            if col_idx == 2:
                c.alignment = Alignment(horizontal='right')
        summary_row += 1

    # Auto width for summary sheet
    for col in ws_summary.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max((len(str(cell.value or '')) for cell in col), default=12)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Remove default empty sheet
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
