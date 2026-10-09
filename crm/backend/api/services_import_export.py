"""CSV and JSON parsing/generation for Services & Garment Items import/export.

Provides:
- export_services_csv(category_ids=None) -> str: CSV text with service items
- export_services_json(category_ids=None) -> list[dict]: structured JSON catalog
- parse_services_rows(uploaded_file) -> (headers, rows) for CSV or XLSX
- import_services_commit(records, overwrite_duplicates=True) -> dict: execution summary
- guess_services_mapping(headers) -> dict: mapping suggestions
"""

import csv
import io
import re

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

from .models import GarmentCategory, GarmentItem, PricingUnit


TARGET_FIELDS = [
    'category',
    'name',
    'price',
    'unit',
    'icon',
    'image_url',
    'is_active',
    'display_order',
    'category_icon',
]

REQUIRED_FIELDS = ['category', 'name', 'price']

_ALIASES = {
    'category': ['category', 'categoryname', 'service', 'servicename', 'section', 'department'],
    'name': ['name', 'item', 'itemname', 'garment', 'garmentname', 'product', 'productname'],
    'price': ['price', 'rate', 'cost', 'unitprice', 'itemprice', 'amount'],
    'unit': ['unit', 'pricingunit', 'uom', 'type'],
    'icon': ['icon', 'itemicon', 'garmenticon'],
    'image_url': ['image', 'imageurl', 'photo', 'picture'],
    'is_active': ['active', 'isactive', 'status', 'enabled'],
    'display_order': ['order', 'displayorder', 'sortorder', 'sequence'],
    'category_icon': ['categoryicon', 'serviceicon'],
}

_UNIT_LOOKUP = {
    'pc': PricingUnit.PIECE,
    'piece': PricingUnit.PIECE,
    'per pc': PricingUnit.PIECE,
    'per piece': PricingUnit.PIECE,
    'kg': PricingUnit.KG,
    'per kg': PricingUnit.KG,
    'sqft': PricingUnit.SQFT,
    'sq.ft': PricingUnit.SQFT,
    'sq ft': PricingUnit.SQFT,
    'per sqft': PricingUnit.SQFT,
    'per sq.ft': PricingUnit.SQFT,
    'set': PricingUnit.SET,
    'per set': PricingUnit.SET,
}


class ImportFileError(ValueError):
    """User-facing error when uploaded file cannot be parsed."""


def _normalize(header):
    return ''.join(ch for ch in str(header).lower() if ch.isalnum())


def guess_services_mapping(headers):
    """Best-effort {target_field: source_header} mapping from column headers."""
    normalized = [(h, _normalize(h)) for h in headers]
    mapping = {}
    for field, aliases in _ALIASES.items():
        for header, norm in normalized:
            if norm in aliases:
                mapping[field] = header
                break
    return mapping


def _cell_to_str(cell):
    if cell is None:
        return ''
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell)


def _sanitize_sheet_title(title, default='Sheet'):
    """Excel sheet names must be <= 31 chars and cannot contain []:*?/\\."""
    clean = re.sub(r'[\[\]:*?/\\]', ' ', str(title or '')).strip()
    return clean[:31] if clean else default


def _parse_xlsx(uploaded_file):
    if openpyxl is None:  # pragma: no cover
        raise ImportFileError('XLSX support is not available on the server.')
    try:
        workbook = openpyxl.load_workbook(uploaded_file, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError(f'Could not read the XLSX file: {exc}') from exc

    # Multi-tab workbook support:
    # Each worksheet represents a Service Category. The items in that sheet belong to that category.
    # If a sheet itself has an explicit 'Category' column (single-sheet flat format), we respect it.
    sheet_data = []
    has_category_col = False

    for sheet in workbook.worksheets:
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            continue
        headers = [_cell_to_str(cell).strip() for cell in header_row]
        if not any(headers):
            continue

        sheet_rows = []
        for row in rows_iter:
            if all(_cell_to_str(cell).strip() == '' for cell in row):
                continue
            sheet_rows.append([_cell_to_str(cell).strip() for cell in row])

        sheet_data.append((sheet.title, headers, sheet_rows))
        if any(_normalize(h) in _ALIASES['category'] for h in headers):
            has_category_col = True

    if not sheet_data:
        return [], []

    # Single sheet with an explicit Category column -> standard flat layout
    if len(sheet_data) == 1 and has_category_col:
        _, headers, rows = sheet_data[0]
        return headers, rows

    # Multi-tab (or single sheet without category column):
    # Synthesize standard layout where 'Category' is derived from sheet.title.
    standard_headers = ['Category', 'Item Name', 'Price', 'Unit', 'Item Icon', 'Image URL', 'Is Active', 'Display Order']
    aggregated_rows = []

    for sheet_title, headers, rows in sheet_data:
        norm_headers = [_normalize(h) for h in headers]
        name_idx = next((i for i, h in enumerate(norm_headers) if h in _ALIASES['name']), None)
        price_idx = next((i for i, h in enumerate(norm_headers) if h in _ALIASES['price']), None)
        unit_idx = next((i for i, h in enumerate(norm_headers) if h in _ALIASES['unit']), None)
        icon_idx = next((i for i, h in enumerate(norm_headers) if h in _ALIASES['icon']), None)
        image_idx = next((i for i, h in enumerate(norm_headers) if h in _ALIASES['image_url']), None)
        active_idx = next((i for i, h in enumerate(norm_headers) if h in _ALIASES['is_active']), None)
        order_idx = next((i for i, h in enumerate(norm_headers) if h in _ALIASES['display_order']), None)
        cat_col_idx = next((i for i, h in enumerate(norm_headers) if h in _ALIASES['category']), None)

        for row in rows:
            cat_val = row[cat_col_idx].strip() if (cat_col_idx is not None and cat_col_idx < len(row) and row[cat_col_idx].strip()) else sheet_title
            item_val = row[name_idx].strip() if (name_idx is not None and name_idx < len(row)) else ''
            price_val = row[price_idx].strip() if (price_idx is not None and price_idx < len(row)) else ''
            unit_val = row[unit_idx].strip() if (unit_idx is not None and unit_idx < len(row)) else ''
            icon_val = row[icon_idx].strip() if (icon_idx is not None and icon_idx < len(row)) else ''
            image_val = row[image_idx].strip() if (image_idx is not None and image_idx < len(row)) else ''
            active_val = row[active_idx].strip() if (active_idx is not None and active_idx < len(row)) else ''
            order_val = row[order_idx].strip() if (order_idx is not None and order_idx < len(row)) else ''

            if item_val or price_val:
                aggregated_rows.append([
                    cat_val,
                    item_val,
                    price_val,
                    unit_val,
                    icon_val,
                    image_val,
                    active_val,
                    order_val,
                ])

    return standard_headers, aggregated_rows


def _parse_csv(uploaded_file):
    raw = uploaded_file.read()
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError as exc:
        raise ImportFileError(f'Could not read the file as text: {exc}') from exc

    reader = csv.reader(io.StringIO(text))
    try:
        headers = [h.strip() for h in next(reader)]
    except StopIteration:
        return [], []

    rows = []
    for row in reader:
        if not any(str(cell).strip() for cell in row):
            continue
        rows.append([str(cell).strip() for cell in row])
    return headers, rows


def parse_services_rows(uploaded_file):
    """Parses XLSX (single or multi-tab) or CSV into (headers, rows)."""
    name = (getattr(uploaded_file, 'name', '') or '').lower()
    if name.endswith('.xlsx'):
        return _parse_xlsx(uploaded_file)
    return _parse_csv(uploaded_file)


def rows_to_records(headers, rows, mapping):
    """Applies {target_field: source_header} mapping to rows."""
    index_by_header = {header: i for i, header in enumerate(headers)}
    records = []
    for row in rows:
        record = {}
        for field, header in mapping.items():
            idx = index_by_header.get(header)
            val = row[idx] if idx is not None and idx < len(row) else ''
            record[field] = val.strip() if isinstance(val, str) else str(val).strip()
        records.append(record)
    return records


def normalize_pricing_unit(raw_unit):
    if not raw_unit:
        return PricingUnit.PIECE
    clean = str(raw_unit).strip().lower()
    return _UNIT_LOOKUP.get(clean, PricingUnit.PIECE)


def parse_bool(value, default=True):
    if value is None or value == '':
        return default
    val_str = str(value).strip().lower()
    if val_str in ('false', '0', 'no', 'inactive', 'off'):
        return False
    if val_str in ('true', '1', 'yes', 'active', 'on'):
        return True
    return default


def parse_float(value, default=0.0):
    if value is None or value == '':
        return default
    try:
        # Strip currency symbols if present (₹, $, commas)
        cleaned = re.sub(r'[^\d.-]', '', str(value))
        return float(cleaned)
    except (ValueError, TypeError):
        return default


def parse_int(value, default=0):
    if value is None or value == '':
        return default
    try:
        return int(float(str(value).strip()))
    except (ValueError, TypeError):
        return default


def export_services_xlsx(category_ids=None):
    """Generates an Excel (.xlsx) workbook where each sheet tab is a Category/Service name
    and the table inside the sheet lists the items and their pricing/details.
    """
    if openpyxl is None:  # pragma: no cover
        raise ImportFileError('XLSX support is not available on the server.')

    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    cat_qs = GarmentCategory.objects.prefetch_related('items').all().order_by(
        'display_order', 'id'
    )
    if category_ids:
        cat_qs = cat_qs.filter(id__in=category_ids)

    wb = openpyxl.Workbook()
    default_sheet = wb.active

    headers = [
        'Item Name',
        'Price',
        'Unit',
        'Item Icon',
        'Image URL',
        'Is Active',
        'Display Order',
    ]

    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='182C4F', end_color='182C4F', fill_type='solid')
    header_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

    cell_font = Font(name='Segoe UI', size=10)
    price_font = Font(name='Segoe UI', size=10, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )

    used_titles = set()

    for idx, category in enumerate(cat_qs):
        base_title = _sanitize_sheet_title(category.name, default=f'Category_{idx+1}')
        sheet_title = base_title
        suffix = 1
        while sheet_title.lower() in used_titles:
            sheet_title = f"{base_title[:28]}_{suffix}"
            suffix += 1
        used_titles.add(sheet_title.lower())

        ws = wb.create_sheet(title=sheet_title)

        # Header row
        ws.row_dimensions[1].height = 26
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # Items
        items = category.items.all().order_by('display_order', 'id')
        for row_idx, item in enumerate(items, 2):
            ws.row_dimensions[row_idx].height = 20
            row_values = [
                item.name,
                float(item.price),
                item.get_unit_display(),
                item.icon,
                item.image_url,
                'Yes' if item.is_active else 'No',
                item.display_order,
            ]
            for col_idx, val in enumerate(row_values, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.font = price_font if col_idx == 2 else cell_font
                c.border = thin_border
                if col_idx == 2:
                    c.number_format = '#,##0.00'
                    c.alignment = Alignment(horizontal='right')
                elif col_idx in (3, 6, 7):
                    c.alignment = Alignment(horizontal='center')

        # Auto-adjust column widths
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    if len(wb.sheetnames) > 1:
        wb.remove(default_sheet)
    else:
        default_sheet.title = 'Services'
        for col_idx, header in enumerate(headers, 1):
            default_sheet.cell(row=1, column=col_idx, value=header)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_services_csv(category_ids=None):
    """Generates a CSV string of all (or selected) services/garment items."""
    qs = GarmentItem.objects.select_related('category').all().order_by(
        'category__display_order', 'display_order', 'id'
    )
    if category_ids:
        qs = qs.filter(category_id__in=category_ids)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        'Category',
        'Category Icon',
        'Item Name',
        'Price',
        'Unit',
        'Item Icon',
        'Image URL',
        'Is Active',
        'Display Order',
    ])

    for item in qs:
        writer.writerow([
            item.category.name,
            item.category.icon,
            item.name,
            f"{item.price:.2f}",
            item.get_unit_display(),
            item.icon,
            item.image_url,
            'Yes' if item.is_active else 'No',
            item.display_order,
        ])

    return output.getvalue()


def export_services_json(category_ids=None):
    """Returns structured JSON/list of categories with nested items."""
    cat_qs = GarmentCategory.objects.prefetch_related('items').all().order_by(
        'display_order', 'id'
    )
    if category_ids:
        cat_qs = cat_qs.filter(id__in=category_ids)

    data = []
    for cat in cat_qs:
        items = []
        for item in cat.items.all().order_by('display_order', 'id'):
            items.append({
                'id': item.id,
                'name': item.name,
                'price': item.price,
                'unit': item.unit,
                'unit_label': item.get_unit_display(),
                'icon': item.icon,
                'image_url': item.image_url,
                'is_active': item.is_active,
                'display_order': item.display_order,
            })
        data.append({
            'id': cat.id,
            'name': cat.name,
            'icon': cat.icon,
            'display_order': cat.display_order,
            'is_active': cat.is_active,
            'items': items,
        })
    return data


def import_services_commit(records, overwrite_duplicates=True):
    """Processes mapped records and creates/updates categories and garment items.

    Returns:
    {
        'total_rows': int,
        'categories_created': int,
        'items_created': int,
        'items_updated': int,
        'skipped_missing': int,
        'skipped_duplicate': int,
    }
    """
    existing_categories = {c.name.strip().lower(): c for c in GarmentCategory.objects.all()}
    existing_items = {
        (item.category_id, item.name.strip().lower()): item
        for item in GarmentItem.objects.select_related('category').all()
    }

    categories_created = 0
    items_created = 0
    items_updated = 0
    skipped_missing = 0
    skipped_duplicate = 0

    to_create_items = []
    to_update_items = []

    for record in records:
        category_name = (record.get('category') or '').strip()
        item_name = (record.get('name') or '').strip()
        raw_price = record.get('price')

        if not category_name or not item_name or raw_price is None or raw_price == '':
            skipped_missing += 1
            continue

        price = parse_float(raw_price)
        unit = normalize_pricing_unit(record.get('unit'))
        icon = (record.get('icon') or '').strip() or 'Shirt'
        image_url = (record.get('image_url') or '').strip()
        is_active = parse_bool(record.get('is_active'), default=True)
        display_order = parse_int(record.get('display_order'), default=0)
        category_icon = (record.get('category_icon') or '').strip() or 'Shirt'

        # Find or create Category
        cat_key = category_name.lower()
        cat = existing_categories.get(cat_key)
        if not cat:
            cat = GarmentCategory.objects.create(
                name=category_name,
                icon=category_icon,
                display_order=len(existing_categories) + 1,
                is_active=True,
            )
            existing_categories[cat_key] = cat
            categories_created += 1

        # Check existing item under this category
        item_key = (cat.id, item_name.lower())
        existing = existing_items.get(item_key)

        if existing:
            if not overwrite_duplicates:
                skipped_duplicate += 1
                continue
            # Update existing
            existing.price = price
            existing.unit = unit
            if icon:
                existing.icon = icon
            if image_url != '':
                existing.image_url = image_url
            existing.is_active = is_active
            if display_order != 0:
                existing.display_order = display_order
            to_update_items.append(existing)
            items_updated += 1
        else:
            new_item = GarmentItem(
                category=cat,
                name=item_name,
                price=price,
                unit=unit,
                icon=icon,
                image_url=image_url,
                is_active=is_active,
                display_order=display_order or (cat.items.count() + len(to_create_items) + 1),
            )
            to_create_items.append(new_item)
            # Store in existing_items in case duplicate occurs within same file
            existing_items[item_key] = new_item
            items_created += 1

    if to_create_items:
        GarmentItem.objects.bulk_create(to_create_items)
    if to_update_items:
        GarmentItem.objects.bulk_update(
            to_update_items,
            ['price', 'unit', 'icon', 'image_url', 'is_active', 'display_order']
        )

    return {
        'total_rows': len(records),
        'categories_created': categories_created,
        'items_created': items_created,
        'items_updated': items_updated,
        'skipped_missing': skipped_missing,
        'skipped_duplicate': skipped_duplicate,
    }
